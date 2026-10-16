"""Leitura de planilhas .xlsx guardadas no Drive.

Nenhum teste fala com o Google: o .xlsx é montado em memória com openpyxl e o
cliente do gspread é trocado por um falso que recusa o arquivo como o Google
recusa ("must not be an Office file").
"""

import datetime as dt
import io

import gspread
import openpyxl
import pytest

from apps.nucleo import sheets


def montar_xlsx():
    livro = openpyxl.Workbook()
    ws = livro.active
    ws.title = "HANDEBOL"
    ws["B2"] = "Custo"
    ws["D2"] = 86.5
    ws["E2"] = 200.0
    ws["B3"] = dt.datetime(2026, 2, 23)
    ws["C3"] = dt.time(21, 30)
    ws["D3"] = True
    ws["E3"] = 7
    ws["B6"] = None  # linha vazia no fim não aparece
    oculta = livro.create_sheet("BASQ")
    oculta["A1"] = "versão antiga"
    oculta.sheet_state = "hidden"
    livro.create_sheet("FUTSAL")["A1"] = "Data"
    dados = io.BytesIO()
    livro.save(dados)
    return dados.getvalue()


class _RespostaFalsa:
    status_code = 400
    text = ""

    def json(self):
        return {"error": {"code": 400, "message": (
            "This operation is not supported for this document. "
            "The document must not be an Office file.")}}


class _ClienteQueRecusaOffice:
    def open_by_key(self, sid):
        raise gspread.exceptions.APIError(_RespostaFalsa())


@pytest.fixture
def xlsx_no_drive(monkeypatch):
    conteudo = montar_xlsx()
    baixados = []

    def _baixar(sid):
        baixados.append(sid)
        return conteudo

    monkeypatch.setattr(sheets, "baixar_arquivo_drive", _baixar)
    monkeypatch.setattr(sheets, "_client", lambda: _ClienteQueRecusaOffice())
    return baixados


def test_grade_sai_no_formato_do_google_em_portugues():
    grade = sheets.abrir_excel("x", conteudo=montar_xlsx()).sheet1.get_all_values()

    assert grade == [
        ["", "", "", "", ""],
        ["", "Custo", "", "86,5", "200"],
        ["", "23/02/2026", "21:30", "TRUE", "7"],
    ]


def test_arquivo_office_cai_no_leitor_de_excel(xlsx_no_drive):
    grade = sheets.ler_grade("https://docs.google.com/spreadsheets/d/ABC123/edit")

    assert xlsx_no_drive == ["ABC123"]
    assert grade[1][1] == "Custo"


def test_aba_por_nome_ignora_maiusculas_e_espacos(xlsx_no_drive):
    grade = sheets.ler_grade("ABC123", aba=" futsal ")

    assert grade == [["Data"]]


def test_ler_abas_deixa_as_ocultas_de_fora_e_baixa_uma_vez(xlsx_no_drive):
    grades = sheets.ler_abas("ABC123")

    assert list(grades) == ["HANDEBOL", "FUTSAL"]
    assert xlsx_no_drive == ["ABC123"]
    assert "BASQ" in sheets.ler_abas("ABC123", incluir_ocultas=True)


def test_ler_aba_com_cabecalho_funciona_em_excel(xlsx_no_drive):
    registros, cabecalho = sheets.ler_aba("ABC123")

    assert cabecalho == ["Custo", "86,5", "200"]
    assert registros == [{"Custo": "23/02/2026", "86,5": "TRUE", "200": "7"}]


def test_arquivo_que_nao_e_excel_vira_erro_claro():
    with pytest.raises(sheets.SheetsIndisponivel, match="não é uma planilha Excel"):
        sheets.abrir_excel("x", conteudo=b"isto nao e um xlsx")


def test_outro_erro_400_do_google_nao_vira_excel(monkeypatch):
    class _Resposta(_RespostaFalsa):
        def json(self):
            return {"error": {"code": 400, "message": "Invalid request"}}

    class _Cliente:
        def open_by_key(self, sid):
            raise gspread.exceptions.APIError(_Resposta())

    monkeypatch.setattr(sheets, "baixar_arquivo_drive", lambda sid: pytest.fail("baixou"))
    with pytest.raises(sheets.SheetsIndisponivel, match="Erro da API"):
        sheets.abrir_planilha(cliente_gs=_Cliente(), sheet_id="ABC123")
